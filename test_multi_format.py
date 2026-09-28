"""Multi-format upload parsing: md/txt/pdf/docx -> plain text via _extract_text."""
import server


def test_extract_md():
    from pathlib import Path
    p = Path(__file__).parent / "_tmp_test.md"
    p.write_text("# 标题\n\n正文内容。", encoding="utf-8")
    try:
        assert "正文内容" in server._extract_text(p)
    finally:
        p.unlink(missing_ok=True)


def test_extract_pdf(tmp_path):
    import fitz
    p = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Computer Architecture Cache Hierarchy")
    doc.save(str(p)); doc.close()
    out = server._extract_text(p)
    assert "Cache Hierarchy" in out or "Computer Architecture" in out


def test_extract_docx(tmp_path):
    from docx import Document
    p = tmp_path / "sample.docx"
    d = Document()
    d.add_heading("操作系统笔记", level=1)
    d.add_paragraph("进程和线程的区别。")
    d.save(str(p))
    out = server._extract_text(p)
    assert "操作系统笔记" in out
    assert "进程和线程" in out


def test_extract_txt(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("数据结构 栈与队列", encoding="utf-8")
    assert "栈与队列" in server._extract_text(p)


def test_unsupported_extension(tmp_path):
    p = tmp_path / "bad.exe"
    p.write_bytes(b"MZ")
    out = server._extract_text(p)
    assert out.startswith("[UNSUPPORTED_FORMAT]")